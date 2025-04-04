import os
import sys
from openpyxl import load_workbook

def main():
    try:
        # Try to load the workbook
        wb = load_workbook('example_input.xlsx', read_only=True)
        
        # Print available sheet names
        print(f"Available sheets: {wb.sheetnames}")
        
        # Try to access the sheet specified in config
        sheet_name = "utrc_active_allocations"
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            # Fall back to the first sheet
            sheet_name = wb.sheetnames[0]
            ws = wb[sheet_name]
            print(f"Sheet 'utrc_active_allocations' not found, using '{sheet_name}' instead")
        
        # Get the header row
        header_row = next(ws.rows)
        headers = [cell.value for cell in header_row]
        print(f"Headers: {headers}")
        
        # Print the first few rows
        print("\nFirst few rows:")
        row_count = 0
        for row in ws.rows:
            if row_count == 0:  # Skip header row
                row_count += 1
                continue
            if row_count > 5:  # Only show 5 rows
                break
            values = [cell.value for cell in row]
            print(values)
            row_count += 1
            
    except Exception as e:
        print(f"Error: {e}")
        return 1
    
    return 0

if __name__ == "__main__":
    sys.exit(main())
